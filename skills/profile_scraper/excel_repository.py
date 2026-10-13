"""
excel_repository.py — Immediate Excel Persistence Layer for 4-Sheet Candidate Lifecycle.

Maintains 4 dedicated sheets:
1. ALL_RAW_PROFILES: Permanent raw extracted profile records & full JSON audit snapshots.
2. ELIGIBLE_CANDIDATES: Fully qualified candidate profiles matching all mandatory requirements.
3. NEEDS_REVIEW: High-potential candidates with unstated public attributes requiring recruiter verification.
4. REJECTED_CANDIDATES: Candidates failing mandatory rules with exact failure reasons.

Implements immediate atomic persistence per candidate with lock recovery and retry handling.
"""

import os
import json
import time
from datetime import datetime
from typing import Set, Tuple, Optional, Any
from loguru import logger
import openpyxl
from openpyxl.utils import get_column_letter

from skills.profile_scraper.candidate_models import (
    CandidateQualificationDecision,
    EligibilityDecision,
    NormalizedCandidateProfile,
    CandidateAnalysisResult
)

# Standardized 4-Sheet Header Columns
COLUMNS_RAW = [
    "Full Name", "Headline", "Location", "Current Role", "Company", "Connection Count",
    "Open to Work", "Total Roles", "Education", "Skills", "Certifications", "Projects", "Languages",
    "Email", "Phone", "Websites", "LinkedIn URL", "GitHub URL",
    "Declared Pronouns", "Explicit DOB", "Scrape Timestamp", "Full Raw JSON"
]

COLUMNS_ELIGIBLE = [
    "Full Name", "Current Role", "Company", "Location", "Total Exp (Yrs)",
    "Experience Bracket", "Relevant Tech Exp (Yrs)", "Primary Domain", "Tech Specialization",
    "Skill Relevance Score", "Top Skills", "Highest Education",
    "Work Authorization / Visa Status", "Open to Work", "Email", "Phone",
    "SignalHire Verified Phone", "SignalHire Verified Email", "SignalHire Status",
    "LinkedIn URL", "GitHub URL", "Date Qualified"
]

COLUMNS_NEEDS_REVIEW = [
    "Full Name", "Current Role", "Company", "Location", "Total Exp (Yrs)",
    "Experience Bracket", "Relevant Tech Exp (Yrs)", "Primary Domain", "Key Missing / Unverified Items",
    "Review Recommendations", "Skill Relevance Score", "Work Authorization / Visa Status",
    "Open to Work", "Email", "Phone",
    "SignalHire Verified Phone", "SignalHire Verified Email", "SignalHire Status",
    "LinkedIn URL", "GitHub URL", "Date Processed"
]

COLUMNS_REJECTED = [
    "Full Name", "Current Role", "Company", "Location", "Total Exp (Yrs)",
    "Experience Bracket", "Primary Domain", "Open to Work", "Rejection Stage", "Primary Rejection Reason",
    "Failed Mandatory Criteria", "LinkedIn URL", "Date Rejected"
]


class ExcelCandidateRepository:
    """
    Manages atomic writing of candidate profiles into the 4-sheet Excel database.
    """

    def __init__(self, excel_path: Optional[str] = None):
        target = excel_path or os.getenv("EXCEL_OUTPUT_PATH", "skills/output/profiles.xlsx")
        if not os.path.isabs(target):
            workspace_root = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
            target = os.path.abspath(os.path.join(workspace_root, target))
        self.excel_path = target
        os.makedirs(os.path.dirname(os.path.abspath(self.excel_path)), exist_ok=True)

    def _get_or_create_sheet(self, wb: openpyxl.Workbook, title: str, columns: list) -> openpyxl.worksheet.worksheet.Worksheet:
        """Retrieves or creates sheet and ensures exact header columns."""
        # Prune legacy sheets if any exist
        for s in list(wb.sheetnames):
            if s in ['Sheet', 'SET_1', 'SET_2', 'SET_3'] or s.startswith('_Rejected_'):
                wb.remove(wb[s])
        if title in wb.sheetnames:
            ws = wb[title]
            first_row = list(ws.iter_rows(max_row=1, values_only=True))
            header = list(first_row[0]) if first_row else []
            if header != columns:
                logger.warning(f"[ExcelRepository] Header mismatch in sheet '{title}'. Recreating sheet.")
                wb.remove(ws)
                ws = wb.create_sheet(title=title)
                ws.append(columns)
        else:
            ws = wb.create_sheet(title=title)
            ws.append(columns)
        return ws

    def _get_existing_urls(self, ws: openpyxl.worksheet.worksheet.Worksheet, url_col_idx: int) -> Set[str]:
        """Collects existing profile URLs in sheet to prevent duplicate entries."""
        urls: Set[str] = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row and len(row) > url_col_idx:
                val = str(row[url_col_idx] or "").strip().lower()
                if val:
                    urls.add(val)
        return urls

    def save_candidate_decision(self, decision: CandidateQualificationDecision) -> bool:
        """
        Immediately persists a processed candidate into the 4-sheet database:
        - ALL_RAW_PROFILES (all candidates)
        - Specific sheet based on decision (ELIGIBLE_CANDIDATES, NEEDS_REVIEW, REJECTED_CANDIDATES)
        """
        c = decision.candidate
        a = decision.analysis
        cand_url = c.linkedin_url.strip().lower()
        max_retries = 5
        retry_delay = 1.0

        wb = None
        for attempt in range(max_retries):
            try:
                if os.path.exists(self.excel_path):
                    wb = openpyxl.load_workbook(self.excel_path)
                else:
                    wb = openpyxl.Workbook()
                # Prune default sheet and any legacy sheets to keep workbook neat
                for s in list(wb.sheetnames):
                    if s in ['Sheet', 'SET_1', 'SET_2', 'SET_3'] or s.startswith('_Rejected_'):
                        wb.remove(wb[s])
                break
            except PermissionError:
                logger.warning(f"[ExcelRepository] File {self.excel_path} locked. Retrying in {retry_delay}s ({attempt+1}/{max_retries})...")
                time.sleep(retry_delay)
                retry_delay *= 2
            except Exception as e:
                logger.error(f"[ExcelRepository] Error opening {self.excel_path}: {e}")
                break

        if wb is None:
            fallback_path = self.excel_path.replace(".xlsx", f"_pending_{int(time.time())}.xlsx")
            logger.error(f"[ExcelRepository] Could not write to main database file. Using fallback: {fallback_path}")
            wb = openpyxl.Workbook()
            if 'Sheet' in wb.sheetnames:
                wb.remove(wb['Sheet'])
            target_path = fallback_path
        else:
            target_path = self.excel_path

        # 1. Write to ALL_RAW_PROFILES
        ws_raw = self._get_or_create_sheet(wb, "ALL_RAW_PROFILES", COLUMNS_RAW)
        raw_urls = self._get_existing_urls(ws_raw, url_col_idx=16)
        otw_str = f"YES ({getattr(c, 'open_to_work_details', 'Verified')})" if getattr(c, "is_open_to_work", False) else "NO"
        if cand_url not in raw_urls:
            raw_json_str = json.dumps(c.raw_payload or {}, default=str)
            if len(raw_json_str) > 32000:
                raw_json_str = raw_json_str[:32000] + "..."

            ws_raw.append([
                c.full_name,
                c.headline,
                c.location,
                c.current_title,
                c.current_company,
                c.connection_count,
                otw_str,
                len(c.experience_history),
                " | ".join([f"{e.degree or 'Degree'} at {e.institution}" for e in c.education]),
                ", ".join(c.skills),
                ", ".join([cert.name for cert in c.certifications]),
                ", ".join([p.title for p in c.projects]),
                ", ".join(c.languages),
                c.email,
                c.phone,
                ", ".join(c.websites),
                c.linkedin_url,
                c.github_url,
                c.pronouns,
                c.date_of_birth,
                c.scraped_timestamp,
                raw_json_str
            ])

        # 2. Write to Category Specific Sheet
        if decision.decision == EligibilityDecision.ELIGIBLE:
            ws_elig = self._get_or_create_sheet(wb, "ELIGIBLE_CANDIDATES", COLUMNS_ELIGIBLE)
            elig_urls = self._get_existing_urls(ws_elig, url_col_idx=19)
            if cand_url not in elig_urls:
                visa_val = a.work_authorization
                if not visa_val or str(visa_val).strip().lower() in ("not_available", "none", "null", "unknown", ""):
                    visa_val = "Not Available"
                ws_elig.append([
                    c.full_name,
                    c.current_title,
                    c.current_company,
                    c.location,
                    a.total_experience_years,
                    getattr(a, "experience_bracket", ">5 years (Senior)" if a.total_experience_years >= 5.0 else "<5 years (Junior/Mid)"),
                    a.relevant_technology_years,
                    a.primary_professional_domain,
                    a.technology_specialization,
                    a.skill_relevance_score,
                    ", ".join(c.skills[:8]),
                    a.highest_education,
                    visa_val,
                    otw_str,
                    c.email,
                    c.phone,
                    decision.signalhire_phone or getattr(c, "signalhire_phone", "Not Available"),
                    decision.signalhire_email or getattr(c, "signalhire_email", "Not Available"),
                    decision.signalhire_status or getattr(c, "signalhire_status", "Not Available"),
                    c.linkedin_url,
                    c.github_url,
                    decision.processed_timestamp
                ])

        elif decision.decision == EligibilityDecision.NEEDS_REVIEW:
            ws_review = self._get_or_create_sheet(wb, "NEEDS_REVIEW", COLUMNS_NEEDS_REVIEW)
            review_urls = self._get_existing_urls(ws_review, url_col_idx=18)
            if cand_url not in review_urls:
                visa_val = a.work_authorization
                if not visa_val or str(visa_val).strip().lower() in ("not_available", "none", "null", "unknown", ""):
                    visa_val = "Not Available"
                ws_review.append([
                    c.full_name,
                    c.current_title,
                    c.current_company,
                    c.location,
                    a.total_experience_years,
                    getattr(a, "experience_bracket", ">5 years (Senior)" if a.total_experience_years >= 5.0 else "<5 years (Junior/Mid)"),
                    a.relevant_technology_years,
                    a.primary_professional_domain,
                    "; ".join(decision.review_reasons),
                    "; ".join(decision.decision_reasons),
                    a.skill_relevance_score,
                    visa_val,
                    otw_str,
                    c.email,
                    c.phone,
                    decision.signalhire_phone or getattr(c, "signalhire_phone", "Not Available"),
                    decision.signalhire_email or getattr(c, "signalhire_email", "Not Available"),
                    decision.signalhire_status or getattr(c, "signalhire_status", "Not Available"),
                    c.linkedin_url,
                    c.github_url,
                    decision.processed_timestamp
                ])

        elif decision.decision == EligibilityDecision.REJECTED:
            ws_rej = self._get_or_create_sheet(wb, "REJECTED_CANDIDATES", COLUMNS_REJECTED)
            rej_urls = self._get_existing_urls(ws_rej, url_col_idx=11)
            if cand_url not in rej_urls:
                primary_reason = decision.failed_mandatory_rules[0] if decision.failed_mandatory_rules else "Failed mandatory criteria"
                ws_rej.append([
                    c.full_name,
                    c.current_title,
                    c.current_company,
                    c.location,
                    a.total_experience_years,
                    getattr(a, "experience_bracket", ">5 years (Senior)" if a.total_experience_years >= 5.0 else "<5 years (Junior/Mid)"),
                    a.primary_professional_domain,
                    otw_str,
                    "Eligibility Rules Evaluation",
                    primary_reason,
                    "; ".join(decision.failed_mandatory_rules),
                    c.linkedin_url,
                    decision.processed_timestamp
                ])

        # Auto-adjust column widths
        for ws in wb.worksheets:
            for col in ws.columns:
                max_len = 0
                for cell in col:
                    val = str(cell.value or '')
                    if len(val) > max_len:
                        max_len = len(val)
                col_letter = get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

        # Save workbook immediately with retry handling
        retry_delay = 0.5
        for attempt in range(2):
            try:
                wb.save(target_path)
                logger.info(f"[ExcelRepository] Successfully persisted '{c.full_name}' -> [{decision.decision.value}] in {target_path}")
                return True
            except PermissionError:
                logger.warning(f"[ExcelRepository] File {target_path} locked by Excel. Retrying in {retry_delay}s...")
                time.sleep(retry_delay)
                retry_delay *= 2
            except Exception as e:
                logger.error(f"[ExcelRepository] Error saving database file: {e}")
                break

        # Fallback save on persistent lock
        try:
            locked_fallback = target_path.replace(".xlsx", f"_LOCKED_{int(time.time())}.xlsx")
            wb.save(locked_fallback)
            logger.info(f"[ExcelRepository] Successfully saved to fallback file {locked_fallback}")
            return True
        except Exception as ex:
            logger.critical(f"[ExcelRepository] CRITICAL: Could not write Excel fallback: {ex}")
            return False
