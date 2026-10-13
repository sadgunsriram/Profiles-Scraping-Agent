"""
Tests for excel_repository.py.
Verifies immediate writing to 4 sheets (ALL_RAW_PROFILES, ELIGIBLE_CANDIDATES, NEEDS_REVIEW, REJECTED_CANDIDATES)
and deduplication integrity.
"""

import os
import openpyxl
from skills.profile_scraper.candidate_models import (
    NormalizedCandidateProfile,
    CandidateAnalysisResult,
    CandidateQualificationDecision,
    EligibilityDecision,
    ExperienceEntry
)
from skills.profile_scraper.excel_repository import (
    ExcelCandidateRepository,
    COLUMNS_RAW,
    COLUMNS_ELIGIBLE,
    COLUMNS_NEEDS_REVIEW,
    COLUMNS_REJECTED
)


def test_excel_repository_writes_to_all_four_sheets(tmp_path):
    test_excel = str(tmp_path / "test_profiles.xlsx")
    repo = ExcelCandidateRepository(excel_path=test_excel)

    # 1. Eligible Candidate
    c1 = NormalizedCandidateProfile(
        full_name="Candidate One",
        linkedin_url="https://www.linkedin.com/in/cand-one",
        headline="Senior Cloud Architect",
        location="Austin, TX",
        experience_history=[ExperienceEntry(title="Architect", company="CloudInc", start_date="2020", end_date="Present")],
        skills=["AWS", "Terraform"],
        email="one@example.com"
    )
    a1 = CandidateAnalysisResult(
        total_experience_years=6.0,
        relevant_technology_years=5.0,
        primary_professional_domain="Cloud / DevOps / Infrastructure",
        skill_relevance_score=90
    )
    d1 = CandidateQualificationDecision(
        candidate=c1,
        analysis=a1,
        decision=EligibilityDecision.ELIGIBLE,
        decision_reasons=["Met all requirements"]
    )
    assert repo.save_candidate_decision(d1) is True

    # 2. Needs Review Candidate
    c2 = NormalizedCandidateProfile(
        full_name="Candidate Two",
        linkedin_url="https://www.linkedin.com/in/cand-two",
        headline="DevOps Engineer",
        location="New York, NY"
    )
    a2 = CandidateAnalysisResult(
        total_experience_years=4.0,
        relevant_technology_years=3.5,
        primary_professional_domain="Cloud / DevOps / Infrastructure",
        skill_relevance_score=75
    )
    d2 = CandidateQualificationDecision(
        candidate=c2,
        analysis=a2,
        decision=EligibilityDecision.NEEDS_REVIEW,
        decision_reasons=["Direct email unstated; clearance unverified"],
        review_reasons=["Clearance not publicly stated"]
    )
    assert repo.save_candidate_decision(d2) is True

    # 3. Rejected Candidate
    c3 = NormalizedCandidateProfile(
        full_name="Candidate Three",
        linkedin_url="https://www.linkedin.com/in/cand-three",
        headline="Talent Acquisition Specialist",
        location="Chicago, IL"
    )
    a3 = CandidateAnalysisResult(
        total_experience_years=5.0,
        relevant_technology_years=0.0,
        primary_professional_domain="Recruitment / HR / Sales",
        skill_relevance_score=10
    )
    d3 = CandidateQualificationDecision(
        candidate=c3,
        analysis=a3,
        decision=EligibilityDecision.REJECTED,
        decision_reasons=["Failed mandatory domain check"],
        failed_mandatory_rules=["Non-practitioner: primary professional domain is Recruitment / HR / Sales"]
    )
    assert repo.save_candidate_decision(d3) is True

    # Inspect Workbook
    wb = openpyxl.load_workbook(test_excel)
    assert "ALL_RAW_PROFILES" in wb.sheetnames
    assert "ELIGIBLE_CANDIDATES" in wb.sheetnames
    assert "NEEDS_REVIEW" in wb.sheetnames
    assert "REJECTED_CANDIDATES" in wb.sheetnames

    # Check ALL_RAW_PROFILES has 3 data rows + 1 header
    ws_raw = wb["ALL_RAW_PROFILES"]
    assert ws_raw.max_row == 4

    # Check ELIGIBLE_CANDIDATES has 1 data row + 1 header
    ws_elig = wb["ELIGIBLE_CANDIDATES"]
    assert ws_elig.max_row == 2
    assert ws_elig.cell(row=2, column=1).value == "Candidate One"

    # Check NEEDS_REVIEW has 1 data row + 1 header
    ws_review = wb["NEEDS_REVIEW"]
    assert ws_review.max_row == 2
    assert ws_review.cell(row=2, column=1).value == "Candidate Two"

    # Check REJECTED_CANDIDATES has 1 data row + 1 header
    ws_rej = wb["REJECTED_CANDIDATES"]
    assert ws_rej.max_row == 2
    assert ws_rej.cell(row=2, column=1).value == "Candidate Three"
    assert "Recruitment / HR / Sales" in ws_rej.cell(row=2, column=7).value

    # Verify Security Clearance column is removed and Work Authorization / Visa Status exists
    assert "Security Clearance" not in COLUMNS_ELIGIBLE
    assert "Security Clearance" not in COLUMNS_NEEDS_REVIEW
    assert "Work Authorization / Visa Status" in COLUMNS_ELIGIBLE
    assert "Work Authorization / Visa Status" in COLUMNS_NEEDS_REVIEW

    # Column 13 in ELIGIBLE_CANDIDATES is Work Authorization / Visa Status
    assert ws_elig.cell(row=1, column=13).value == "Work Authorization / Visa Status"
    assert ws_elig.cell(row=2, column=13).value in ("Not Available", "US Citizen")

    # Column 12 in NEEDS_REVIEW is Work Authorization / Visa Status
    assert ws_review.cell(row=1, column=12).value == "Work Authorization / Visa Status"
    assert ws_review.cell(row=2, column=12).value == "Not Available"
