"""
Unit tests for Excel output and multi-sheet profile export.
"""

import os
import openpyxl
from skills.profile_scraper.excel_manager import batch_save_profiles
from skills.profile_scraper.set1_identity import Set1Candidate


def test_batch_save_profiles_creates_valid_workbook(tmp_path):
    output_file = str(tmp_path / "test_profiles.xlsx")

    # Sample candidate
    cand1 = Set1Candidate(
        technology="Network Engineer",
        github_url="https://github.com/cand1",
        linkedin_url="https://linkedin.com/in/cand1",
        joined_linkedin=None,
        full_name="Jane Doe",
        years_of_experience=4.5,
        experience_bracket="3-5",
        location="Dallas, TX",
        current_company="Cisco",
        email="jane.doe@example.com",
        phone="+1 214-555-0100",
        skills="BGP, OSPF, Routing"
    )

    rejections = {
        "Location Validation": [
            {
                "full_name": "Bob Reject",
                "technology": "Network",
                "years_of_experience": 2.0,
                "linkedin_url": "https://linkedin.com/in/bobreject",
                "location": "London, UK",
                "rejection_stage": "Location Validation",
                "rejection_reason": "Outside USA"
            }
        ]
    }

    batch_save_profiles(
        set1_candidates=[cand1],
        set2_candidates=[cand1],
        set3_candidates=[cand1],
        rejected_candidates=rejections,
        path=output_file
    )

    assert os.path.exists(output_file)

    wb = openpyxl.load_workbook(output_file)
    assert "ELIGIBLE_CANDIDATES" in wb.sheetnames
    assert "NEEDS_REVIEW" in wb.sheetnames
    assert "REJECTED_CANDIDATES" in wb.sheetnames
    assert "ALL_RAW_PROFILES" in wb.sheetnames
    assert "SET_1" not in wb.sheetnames
    assert "SET_2" not in wb.sheetnames
    assert "SET_3" not in wb.sheetnames
    assert "_Rejected_Location Validation" not in wb.sheetnames

    # Check ELIGIBLE_CANDIDATES sheet
    ws_elig = wb["ELIGIBLE_CANDIDATES"]
    rows_elig = list(ws_elig.iter_rows(values_only=True))
    assert rows_elig[1][0] == "Jane Doe"
    assert "Email" in rows_elig[0]
    assert "Phone" in rows_elig[0]
    assert "Work Authorization / Visa Status" in rows_elig[0]
    email_idx = rows_elig[0].index("Email")
    phone_idx = rows_elig[0].index("Phone")
    assert rows_elig[1][email_idx] == "jane.doe@example.com"
    assert rows_elig[1][phone_idx] == "+1 214-555-0100"

    # Check REJECTED_CANDIDATES sheet
    ws_rej = wb["REJECTED_CANDIDATES"]
    rows_rej = list(ws_rej.iter_rows(values_only=True))
    assert rows_rej[1][0] == "Bob Reject"
    assert rows_rej[1][8] == "Location Validation"
    assert "Outside USA" in rows_rej[1][9]

    wb.close()
