import os
import sys
import openpyxl
from bs4 import BeautifulSoup

# Ensure project root is in path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from skills.profile_scraper.profile_enrichment import (
    detect_missing_fields,
    calculate_identity_confidence,
    extract_explicit_clearance,
    extract_explicit_visa,
    extract_contact_info_from_page,
    enrich_candidate_multisource,
    EvidenceLedger,
    EvidenceItem
)
from skills.profile_scraper.set1_identity import Set1Candidate, qualify_candidate_set1
import skills.profile_scraper.set2_clearance as set2_clearance
import skills.profile_scraper.set3_contact as set3_contact
import skills.profile_scraper.excel_manager as excel_manager


def test_missing_field_detection():
    print("\n--- 1. Testing Missing Field Detection ---")
    p_full = {
        "email": "alex@company.com",
        "mobile_phone": "+12065550199",
        "visa_status": "US Citizen",
        "security_clearance": "Secret",
        "github_url": "https://github.com/alexsmith",
        "website": "https://alexsmith.dev",
        "skills": ["Python", "AWS", "Docker"]
    }
    missing_full = detect_missing_fields(p_full)
    print("Full profile missing fields:", missing_full)
    assert not any(missing_full.values()), f"Expected no missing fields, got {missing_full}"

    p_sparse = {
        "email": "",
        "mobile_phone": "",
        "visa_status": "Unknown",
        "security_clearance": "Not available",
        "github_url": "",
        "website": "",
        "skills": []
    }
    missing_sparse = detect_missing_fields(p_sparse)
    print("Sparse profile missing fields:", missing_sparse)
    assert missing_sparse["email"] is True
    assert missing_sparse["phone"] is True
    assert missing_sparse["visa_status"] is True
    assert missing_sparse["security_clearance"] is True
    assert missing_sparse["github_url"] is True
    assert missing_sparse["portfolio"] is True
    assert missing_sparse["skills"] is True
    print("Missing field detection tests passed!")


def test_identity_confidence_scoring():
    print("\n--- 2. Testing Identity Confidence Scoring Matrix ---")
    candidate = {
        "full_name": "Marcus Vance",
        "current_company": "Northrop Grumman",
        "location": "Dallas, TX",
        "technology": "Network Engineer",
        "github_url": "https://github.com/marcusvance",
        "linkedin_url": "https://linkedin.com/in/marcus-vance-net"
    }

    # Match A: High confidence matching page (Name + Company + Location + Tech)
    page_title_a = "Marcus Vance - Senior Network Engineer"
    page_text_a = "Marcus Vance is a Network Engineer based in Dallas, TX currently working at Northrop Grumman."
    score_a, sigs_a, is_ver_a = calculate_identity_confidence(candidate, page_title_a, page_text_a, "https://marcusvance.dev")
    print(f"Match A Score: {score_a}/100, Verified: {is_ver_a}, Signals: {sigs_a}")
    assert score_a >= 65, f"Expected verified score >= 65, got {score_a}"
    assert is_ver_a is True

    # Match B: Name collision on wrong person (Different company, different location, different field)
    page_title_b = "Marcus Vance - Real Estate Agent in Miami"
    page_text_b = "Marcus Vance is a premier real estate agent helping buyers in Miami, Florida."
    score_b, sigs_b, is_ver_b = calculate_identity_confidence(candidate, page_title_b, page_text_b, "https://marcusvancerealty.com")
    print(f"Match B Score: {score_b}/100, Verified: {is_ver_b}, Signals: {sigs_b}")
    assert score_b < 65, f"Expected non-verified score < 65 for collision, got {score_b}"
    assert is_ver_b is False

    # Match C: GitHub handle match in text
    page_title_c = "Personal Portfolio"
    page_text_c = "Welcome to my portfolio! Check out my projects on github.com/marcusvance or reach out."
    score_c, sigs_c, is_ver_c = calculate_identity_confidence(candidate, page_title_c, page_text_c, "https://marcus.dev")
    print(f"Match C Score: {score_c}/100, Verified: {is_ver_c}, Signals: {sigs_c}")
    assert "GitHub Handle Match" in sigs_c
    print("Identity scoring tests passed!")


def test_clearance_and_visa_extraction():
    print("\n--- 3. Testing Explicit Clearance & Visa Extraction ---")
    
    # Clearance extraction
    text_clr = "Senior Systems Engineer with an active TS/SCI security clearance and DoD experience."
    clr, snip = extract_explicit_clearance(text_clr)
    print(f"Clearance: {clr}, Snippet: '{snip}'")
    assert clr == "TS/SCI"

    text_no_clr = "Software developer working at defense contractor Lockheed Martin."
    clr_none, _ = extract_explicit_clearance(text_no_clr)
    print(f"Clearance (no explicit mention): {clr_none}")
    assert clr_none is None, "Should not infer clearance from employer name"

    # Visa extraction
    text_visa = "Experienced cloud architect. U.S. Citizen, no sponsorship required."
    v_stat, v_snip = extract_explicit_visa(text_visa)
    print(f"Visa: {v_stat}, Snippet: '{v_snip}'")
    assert v_stat == "US Citizen"

    text_no_visa = "Software engineer from India living in Texas."
    v_none, _ = extract_explicit_visa(text_no_visa)
    print(f"Visa (no explicit status): {v_none}")
    assert v_none is None, "Should not infer visa from nationality or location"
    print("Clearance and visa extraction tests passed!")


def test_full_pipeline_with_enrichment_and_excel():
    print("\n--- 4. Testing Pipeline with Evidence Ledger & Excel Export ---")
    
    mock_profile = {
        "full_name": "Jordan Lee",
        "location": "Seattle, Washington, USA",
        "github_url": "https://github.com/jordanlee",
        "linkedin_url": "https://linkedin.com/in/jordanlee",
        "technology": "Cloud Engineer",
        "headline": "Lead Cloud Infrastructure Engineer",
        "bio": "Cloud Engineer at Microsoft with active Secret clearance. US Citizen.",
        "skills": ["Terraform", "Kubernetes", "AWS", "Python"],
        "experience_history": [
            {"title": "Lead Cloud Infrastructure Engineer", "company": "Microsoft", "start_date": "Jan 2020", "end_date": "Present"}
        ],
        "blog": "https://jordanlee.dev"
    }

    # Run multi-source enrichment on mock profile
    enriched_profile, ledger = enrich_candidate_multisource(mock_profile, "Cloud Engineer")
    print(f"Enrichment Completed for: {enriched_profile['full_name']}")
    print(f"  Clearance: {enriched_profile.get('security_clearance')}")
    print(f"  Visa: {enriched_profile.get('visa_status')}")
    print(f"  Evidence Ledger Keys: {list(enriched_profile.get('evidence_ledger', {}).keys())}")

    # Pass to SET 1
    candidate = qualify_candidate_set1(enriched_profile, "Cloud Engineer")
    assert candidate is not None, "Candidate must pass SET 1"
    print(f"SET 1 Passed: {candidate.full_name}, Exp: {candidate.years_of_experience} yrs ({candidate.experience_bracket})")

    # Pass to SET 2
    passed_s2, rej_stage, rej_reason, visa_spons, clearance = set2_clearance.qualify_candidate_set2(candidate, enriched_profile, "Cloud Engineer")
    assert passed_s2 is True, f"SET 2 Failed: {rej_reason}"
    candidate.visa_sponsorship_needed = visa_spons
    candidate.security_clearance = clearance
    candidate.visa_status = enriched_profile.get("visa_status", "Unknown")
    candidate.current_pipeline_stage = "Set 2 Passed"
    print(f"SET 2 Passed: Clearance='{clearance}', Visa='{candidate.visa_status}'")

    # Pass to SET 3
    passed_s3, rej_s3, rej_r3, v_email, f_phone, website, portfolio = set3_contact.qualify_candidate_set3(candidate, enriched_profile)
    assert passed_s3 is True, f"SET 3 Failed: {rej_r3}"
    candidate.email = v_email
    candidate.phone = f_phone
    candidate.website = website
    candidate.portfolio = portfolio
    candidate.current_pipeline_stage = "Qualified"
    print(f"SET 3 Passed: Email='{v_email}', Phone='{f_phone}', Portfolio='{portfolio}'")

    # Verify Excel export
    test_excel_path = "scratch/test_enriched_profiles.xlsx"
    if os.path.exists(test_excel_path):
        os.remove(test_excel_path)

    excel_manager.batch_save_profiles([candidate], [candidate], [candidate], {"Set1": [], "Set2": [], "Set3": []}, test_excel_path)
    
    wb = openpyxl.load_workbook(test_excel_path)
    print("Created Sheets in Excel:", wb.sheetnames)
    assert "SET_1" in wb.sheetnames
    assert "SET_2" in wb.sheetnames
    assert "SET_3" in wb.sheetnames
    
    ws_set3 = wb["SET_3"]
    rows = list(ws_set3.iter_rows(values_only=True))
    headers = rows[0]
    data_row = rows[1]
    print(f"SET_3 Headers: {headers}")
    print(f"SET_3 Row 1  : {data_row}")
    
    assert "Clearance Source" in headers
    assert "Visa Evidence" in headers
    assert "Email Verified" in headers
    assert "Portfolio URL" in headers
    assert data_row[0] == "Jordan Lee"
    print("Pipeline and Excel verification tests passed!")


if __name__ == "__main__":
    test_missing_field_detection()
    test_identity_confidence_scoring()
    test_clearance_and_visa_extraction()
    test_full_pipeline_with_enrichment_and_excel()
    print("\n=======================================================")
    print("  ALL MULTI-SOURCE ENRICHMENT TESTS PASSED SUCCESSFULLY! ")
    print("=======================================================\n")
