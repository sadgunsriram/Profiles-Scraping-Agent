import openpyxl
import os

excel_path = "skills/output/profiles.xlsx"
if not os.path.exists(excel_path):
    # Fallback to local skills/profile_scraper/profiles.xlsx if output is elsewhere
    excel_path = "skills/profile_scraper/profiles.xlsx"

print(f"Loading Excel file: {excel_path}...")
try:
    wb = openpyxl.load_workbook(excel_path)
    print("Sheets available:", wb.sheetnames)
    
    # Let's inspect the first non-system sheet
    target_sheet = None
    for name in wb.sheetnames:
        if not name.startswith("_"):
            target_sheet = name
            break
            
    if target_sheet:
        ws = wb[target_sheet]
        print(f"\nInspecting sheet: '{target_sheet}'")
        
        # Read header row (Row 1)
        headers = [cell.value for cell in ws[1]]
        print(f"Total Columns ({len(headers)}):", headers)
        
        # Read first data row (Row 2)
        if ws.max_row >= 2:
            first_row = [cell.value for cell in ws[2]]
            print("\nSample Data (Row 2):")
            for col, val in zip(headers, first_row):
                print(f"  - {col:19s}: {val}")
        else:
            print("No data rows found in this sheet yet.")
    else:
        print("No technology sheets found in the workbook yet.")
        
except Exception as e:
    print(f"Error reading workbook: {e}")
