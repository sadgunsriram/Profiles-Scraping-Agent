import openpyxl
import os

def check_file(path):
    print(f"\n========================================\nFile: {path}")
    if not os.path.exists(path):
        print("File does not exist")
        return
        
    wb = openpyxl.load_workbook(path)
    required_columns = ["Name", "Technology", "Years of exp", "LinkedIn", "GitHub",
                        "Email", "Location", "Security Clearances", "Skill set"]
                        
    for sheet_name in wb.sheetnames:
        if sheet_name.endswith('_Incomplete') or sheet_name in ('Profiles', 'Incomplete', 'Sheet', '_Rejected', 'Other', 'Test'):
            continue
        ws = wb[sheet_name]
        header = [cell.value for cell in ws[1]]
        
        try:
            col_indices = {col: header.index(col) for col in required_columns if col in header}
        except ValueError as e:
            print(f"Error finding columns in sheet {sheet_name}: {e}")
            continue
            
        strict_count = 0
        logical_count = 0
        
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            # Check strict completeness
            strict_ok = True
            for col in required_columns:
                if col not in col_indices:
                    strict_ok = False
                    break
                idx = col_indices[col]
                val = row[idx] if idx < len(row) else None
                if val is None or str(val).strip() == '':
                    strict_ok = False
                    break
            if strict_ok:
                strict_count += 1
                
            # Check logical completeness
            name_idx = col_indices.get("Name")
            loc_idx = col_indices.get("Location")
            email_idx = col_indices.get("Email")
            li_idx = col_indices.get("LinkedIn")
            gh_idx = col_indices.get("GitHub")
            
            name = row[name_idx] if name_idx is not None and name_idx < len(row) else None
            loc = row[loc_idx] if loc_idx is not None and loc_idx < len(row) else None
            email = row[email_idx] if email_idx is not None and email_idx < len(row) else None
            li = row[li_idx] if li_idx is not None and li_idx < len(row) else None
            gh = row[gh_idx] if gh_idx is not None and gh_idx < len(row) else None
            
            has_contact = (email and str(email).strip()) or (li and str(li).strip()) or (gh and str(gh).strip())
            if name and str(name).strip() and loc and str(loc).strip() and has_contact:
                logical_count += 1
                
        print(f"Sheet: {sheet_name}")
        print(f"  Total data rows: {ws.max_row - 1}")
        print(f"  Strictly complete: {strict_count}")
        print(f"  Logically complete: {logical_count}")

print("Checking backup file:")
check_file("skills/output/profiles_backup.xlsx")
print("\nChecking current output file:")
check_file("skills/output/profiles.xlsx")
