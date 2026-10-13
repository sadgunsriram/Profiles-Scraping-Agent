import os
import sys
import openpyxl

# Ensure project root is on path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from skills.profile_scraper.profile_enrichment import (
    inspect_and_extract_linkedin_fields,
    extract_explicit_clearance_from_text,
    extract_explicit_visa_from_text,
    extract_contacts_from_linkedin_text
)
from skills.profile_scraper.set1_identity import Set1Candidate, qualify_candidate_set1
import skills.profile_scraper.set2_clearance as set2_clearance
import skills.profile_scraper.set3_contact as set3_contact
import skills.profile_scraper.excel_manager as excel_manager


def run_pipeline_for_candidate(raw_profile: dict, query: str = "Cloud Engineer"):
    # Step 9.5: Inspect & extract from LinkedIn only
    inspected_profile, field_status = inspect_and_extract_linkedin_fields(raw_profile)
    
    # Step 10: SET 1
    candidate = qualify_candidate_set1(inspected_profile, query)
    assert candidate is not None, f"Candidate {raw_profile.get('full_name')} should pass SET 1"
    
    # Step 11: SET 2 (LinkedIn-only, never rejected for missing optional fields)
    passed_s2, rej_stage, rej_reason, visa_spons, clearance = set2_clearance.qualify_candidate_set2(candidate, inspected_profile, query)
    assert passed_s2 is True, f"Candidate should pass SET 2: {rej_reason}"
    candidate.visa_sponsorship_needed = visa_spons
    candidate.security_clearance = clearance
    candidate.visa_status = inspected_profile.get("visa_status", "Not Available")
    candidate.current_pipeline_stage = "Set 2 Passed"
    
    # Step 12: SET 3 (LinkedIn-only, never rejected for missing email/phone)
    passed_s3, rej_s3, rej_r3, v_email, f_phone, website, portfolio = set3_contact.qualify_candidate_set3(candidate, inspected_profile)
    assert passed_s3 is True, f"Candidate should pass SET 3: {rej_r3}"
    candidate.email = v_email
    candidate.phone = f_phone
    candidate.website = website
    candidate.portfolio = portfolio
    candidate.current_pipeline_stage = "Qualified"
    
    return candidate, inspected_profile, field_status


def test_1_all_details_available():
    print("\n--- Test 1: All Details Available on LinkedIn ---")
    profile = {
        "full_name": "Alice Johnson",
        "location": "Dallas, TX, USA",
        "github_url": "https://github.com/alicejohnson",
        "linkedin_url": "https://linkedin.com/in/alicejohnson",
        "technology": "Cloud Engineer",
        "headline": "Senior Cloud Engineer | Active TS/SCI Clearance | H1B Visa",
        "bio": "Cloud infrastructure engineer with TS/SCI clearance and H1B visa. Reach me at alice.j@company.com or +1 214 555 0199.",
        "skills": ["AWS", "Terraform", "Kubernetes"],
        "experience_history": [
            {"title": "Senior Cloud Engineer", "company": "Tech Corp", "start_date": "Jan 2019", "end_date": "Present", "description": "Cloud infra and TS/SCI workflows"}
        ],
        "contact_info": {
            "email": "alice.j@company.com",
            "phone": "+1 214 555 0199"
        }
    }
    cand, _, status = run_pipeline_for_candidate(profile)
    print(f"SET 1: Name='{cand.full_name}', Exp='{cand.years_of_experience}'")
    print(f"SET 2: Visa='{cand.visa_status}', Clearance='{cand.security_clearance}'")
    print(f"SET 3: Email='{cand.email}', Phone='{cand.phone}'")
    
    assert cand.visa_status == "H1B"
    assert cand.security_clearance == "TS/SCI"
    assert cand.email == "alice.j@company.com"
    assert "+12145550199" in cand.phone or "+1 214 555 0199" in cand.phone
    print("Test 1 Passed!")


def test_2_visa_available_clearance_unavailable():
    print("\n--- Test 2: Visa Available, Clearance Unavailable on LinkedIn ---")
    profile = {
        "full_name": "Bob Smith",
        "location": "Austin, TX, USA",
        "github_url": "https://github.com/bobsmith",
        "linkedin_url": "https://linkedin.com/in/bobsmith",
        "technology": "Cloud Engineer",
        "headline": "Cloud Architect | US Citizen",
        "bio": "Cloud specialist. US Citizen, no sponsorship needed.",
        "skills": ["AWS", "Python", "Docker"],
        "experience_history": [
            {"title": "Cloud Architect", "company": "CloudTech", "start_date": "Jan 2018", "end_date": "Present", "description": "Building cloud native apps"}
        ]
    }
    cand, _, status = run_pipeline_for_candidate(profile)
    print(f"SET 2: Visa='{cand.visa_status}', Clearance='{cand.security_clearance}'")
    
    assert cand.visa_status == "US Citizen"
    assert cand.security_clearance == "Not Available"
    assert cand.current_pipeline_stage == "Qualified"
    print("Test 2 Passed!")


def test_3_clearance_available_visa_unavailable():
    print("\n--- Test 3: Clearance Available, Visa Unavailable on LinkedIn ---")
    profile = {
        "full_name": "Charlie Davis",
        "location": "Reston, VA, USA",
        "github_url": "https://github.com/charliedavis",
        "linkedin_url": "https://linkedin.com/in/charliedavis",
        "technology": "Cloud Engineer",
        "headline": "Lead Systems Engineer with Top Secret Clearance",
        "bio": "10+ years engineering defense applications. Active Top Secret.",
        "skills": ["Linux", "Kubernetes", "Security"],
        "experience_history": [
            {"title": "Lead Systems Engineer", "company": "Federal Corp", "start_date": "Jan 2017", "end_date": "Present", "description": "Top Secret classified environments"}
        ]
    }
    cand, _, status = run_pipeline_for_candidate(profile)
    print(f"SET 2: Visa='{cand.visa_status}', Clearance='{cand.security_clearance}'")
    
    assert cand.visa_status == "Not Available"
    assert cand.security_clearance == "TOP SECRET"
    assert cand.current_pipeline_stage == "Qualified"
    print("Test 3 Passed!")


def test_4_both_set2_fields_unavailable():
    print("\n--- Test 4: Both Visa and Clearance Unavailable on LinkedIn ---")
    profile = {
        "full_name": "David Wilson",
        "location": "Seattle, WA, USA",
        "github_url": "https://github.com/davidwilson",
        "linkedin_url": "https://linkedin.com/in/davidwilson",
        "technology": "Cloud Engineer",
        "headline": "Software Engineer at Amazon",
        "bio": "Building scalable backend services.",
        "skills": ["Java", "AWS", "DynamoDB"],
        "experience_history": [
            {"title": "Software Engineer", "company": "Amazon", "start_date": "Jan 2020", "end_date": "Present", "description": "Backend scaling"}
        ]
    }
    cand, _, status = run_pipeline_for_candidate(profile)
    print(f"SET 2: Visa='{cand.visa_status}', Clearance='{cand.security_clearance}'")
    
    assert cand.visa_status == "Not Available"
    assert cand.security_clearance == "Not Available"
    assert cand.current_pipeline_stage == "Qualified"
    print("Test 4 Passed!")


def test_5_email_and_phone_unavailable():
    print("\n--- Test 5: Email and Phone Unavailable on LinkedIn ---")
    profile = {
        "full_name": "Emma Watson",
        "location": "San Francisco, CA, USA",
        "github_url": "https://github.com/emmawatson",
        "linkedin_url": "https://linkedin.com/in/emmawatson",
        "technology": "Cloud Engineer",
        "headline": "DevOps Engineer at Stripe",
        "bio": "Passionate about CI/CD pipelines and Kubernetes infrastructure.",
        "skills": ["Kubernetes", "CI/CD", "Go"],
        "experience_history": [
            {"title": "DevOps Engineer", "company": "Stripe", "start_date": "Jan 2021", "end_date": "Present", "description": "CI/CD automation"}
        ]
    }
    cand, _, status = run_pipeline_for_candidate(profile)
    print(f"SET 3: Email='{cand.email}', Phone='{cand.phone}'")
    
    assert cand.email == "Not Available"
    assert cand.phone == "Not Available"
    assert cand.current_pipeline_stage == "Qualified"
    print("Test 5 Passed! Candidate was saved to SET 3 without rejection.")


def test_6_verify_linkedin_inspection_status():
    print("\n--- Test 6: Verify LinkedIn Inspection Status Tracking ---")
    profile = {
        "full_name": "Frank Miller",
        "location": "Denver, CO, USA",
        "github_url": "https://github.com/frankmiller",
        "linkedin_url": "https://linkedin.com/in/frankmiller",
        "technology": "Cloud Engineer",
        "headline": "Cloud Specialist",
        "bio": "Working on GCP migrations.",
        "skills": ["GCP", "Terraform"],
        "experience_history": [
            {"title": "Cloud Specialist", "company": "DenverTech", "start_date": "Jan 2019", "end_date": "Present", "description": "Cloud migration"}
        ]
    }
    cand, inspected, status = run_pipeline_for_candidate(profile)
    print("Inspection status dict:", status)
    
    assert status["visa_status"]["checked"] is True
    assert status["visa_status"]["found"] is False
    assert status["visa_status"]["value"] == "Not Available"
    
    assert status["security_clearance"]["checked"] is True
    assert status["security_clearance"]["found"] is False
    assert status["security_clearance"]["value"] == "Not Available"
    
    assert status["email"]["checked"] is True
    assert status["email"]["found"] is False
    assert status["email"]["value"] == "Not Available"
    
    assert status["phone"]["checked"] is True
    assert status["phone"]["found"] is False
    assert status["phone"]["value"] == "Not Available"
    print("Test 6 Passed! System explicitly checked LinkedIn sections before assigning 'Not Available'.")


def test_7_excel_storage_and_schema():
    print("\n--- Test 7: Excel Export with Standardized Schemas ---")
    candidates = []
    
    profiles = [
        {"full_name": "Candidate One", "headline": "TS/SCI cleared", "location": "USA", "skills": ["AWS"], "experience_history": [{"title": "Eng", "company": "A", "start_date": "Jan 2020", "end_date": "Present"}]},
        {"full_name": "Candidate Two", "headline": "H1B Engineer", "location": "USA", "skills": ["AWS"], "experience_history": [{"title": "Eng", "company": "B", "start_date": "Jan 2020", "end_date": "Present"}]},
        {"full_name": "Candidate Three", "headline": "DevOps", "location": "USA", "skills": ["AWS"], "experience_history": [{"title": "Eng", "company": "C", "start_date": "Jan 2020", "end_date": "Present"}]}
    ]
    for p in profiles:
        p["github_url"] = f"https://github.com/{p['full_name'].replace(' ', '').lower()}"
        p["linkedin_url"] = f"https://linkedin.com/in/{p['full_name'].replace(' ', '').lower()}"
        cand, _, _ = run_pipeline_for_candidate(p)
        candidates.append(cand)
        
    excel_path = "scratch/test_linkedin_only_profiles.xlsx"
    if os.path.exists(excel_path):
        os.remove(excel_path)
        
    excel_manager.batch_save_profiles(candidates, candidates, candidates, {}, excel_path)
    
    wb = openpyxl.load_workbook(excel_path)
    print("Sheets created:", wb.sheetnames)
    assert wb.sheetnames == ["SET_1", "SET_2", "SET_3"]
    
    # Check SET_1
    ws1 = wb["SET_1"]
    row1_h = list(ws1.iter_rows(values_only=True))[0]
    print("SET_1 Headers:", row1_h)
    assert list(row1_h) == ["Full Name", "Technology", "Years of Experience", "LinkedIn URL", "GitHub URL"]
    
    # Check SET_2
    ws2 = wb["SET_2"]
    row2_h = list(ws2.iter_rows(values_only=True))[0]
    print("SET_2 Headers:", row2_h)
    assert list(row2_h) == ["Full Name", "Technology", "Years of Experience", "LinkedIn URL", "GitHub URL", "Location", "Security Clearances", "Visa Status"]
    
    # Check SET_3
    ws3 = wb["SET_3"]
    row3_h = list(ws3.iter_rows(values_only=True))[0]
    row3_d = list(ws3.iter_rows(values_only=True))[1:]
    print("SET_3 Headers:", row3_h)
    assert list(row3_h) == ["Full Name", "Technology", "Years of Experience", "LinkedIn URL", "GitHub URL", "Location", "Security Clearances", "Visa Status", "Email", "Phone", "Skill set"]
    
    for row in row3_d:
        print("  SET_3 Row:", row)
        # Ensure no None or empty string in optional fields (indices 6, 7, 8, 9)
        assert row[6] in ("TS/SCI", "Not Available")
        assert row[7] in ("H1B", "Not Available")
        assert row[8] == "Not Available"
        assert row[9] == "Not Available"
        
    print("Test 7 Passed! Excel output verified.")


if __name__ == "__main__":
    test_1_all_details_available()
    test_2_visa_available_clearance_unavailable()
    test_3_clearance_available_visa_unavailable()
    test_4_both_set2_fields_unavailable()
    test_5_email_and_phone_unavailable()
    test_6_verify_linkedin_inspection_status()
    test_7_excel_storage_and_schema()
    print("\n=======================================================")
    print("  ALL LINKEDIN-ONLY PIPELINE TESTS PASSED 100%!        ")
    print("=======================================================\n")
