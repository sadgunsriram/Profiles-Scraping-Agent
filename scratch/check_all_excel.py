import openpyxl
import os
import glob

def check_file(path):
    print(f"\n========================================\nFile: {path} ({os.path.getsize(path)} bytes)")
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        print(f"Sheets: {wb.sheetnames}")
        for sheet in wb.sheetnames:
            try:
                ws = wb[sheet]
                # ws.max_row might be slow or inaccurate in read_only, but let's try to get a quick estimate
                row_count = 0
                for _ in ws.iter_rows(values_only=True):
                    row_count += 1
                print(f"  - {sheet}: {row_count} rows (including header)")
            except Exception as e:
                print(f"  - {sheet}: Error counting rows: {e}")
        wb.close()
    except Exception as e:
        print(f"Error reading file: {e}")

# Find all xlsx files
files = glob.glob("**/*.xlsx", recursive=True)
for file in sorted(files):
    if "~$" in file:
        continue
    check_file(file)
