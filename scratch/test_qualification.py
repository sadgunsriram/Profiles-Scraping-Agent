import os
import sys
import openpyxl

# Add project root to sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from skills.profile_scraper.set1_identity import Set1Candidate
import skills.profile_scraper.set2_clearance as set2_clearance
import skills.profile_scraper.set3_contact as set3_contact
import skills.profile_scraper.excel_manager as excel_manager

def run_tests():
    print("=== Testing SET 2 & SET 3 Qualification Logic ===")
    
    # 1. Mock US candidate without clearance or explicit visa
    candidate = Set1Candidate(
        technology="Network Engineer",
        github_url="https://github.com/testuser",
        linkedin_url="https://linkedin.com/in/testuser",
        joined_linkedin=2020,
        full_name="Alex Smith",
        years_of_experience=6,
        experience_bracket="Senior",
        location="Austin, Texas, USA",
        current_company="Acme Corp",
        email="",
        phone="",
        confidence_score="High",
        experience_validation_status="Valid",
        current_pipeline_stage="Set 1 Passed",
        skills="Cisco, Routing, BGP, Python"
    )
    
    profile = {
        "full_name": "Alex Smith",
        "location": "Austin, Texas, USA",
        "bio": "Network Engineer at Acme Corp. Love Cisco and BGP.",
        "headline": "Senior Network Engineer",
        "skills": ["Cisco", "Routing", "BGP", "Python"],
        "experience_history": [{"title": "Senior Network Engineer", "company": "Acme Corp"}]
    }

    # Test A: General Search ("Network Engineer")
    query_general = "Network Engineer"
    passed_s2, rej_stage, rej_reason, visa_spons, clearance = set2_clearance.qualify_candidate_set2(candidate, profile, query_general)
    print(f"Test A (General Query Set 2): Passed={passed_s2}, Clearance='{clearance}', Visa='{profile.get('visa_status')}'")
    assert passed_s2 is True, f"Expected True for general query Set 2, got {passed_s2} ({rej_reason})"
    
    # Update candidate fields from Set 2
    candidate.visa_sponsorship_needed = visa_spons
    candidate.security_clearance = clearance
    candidate.visa_status = profile.get("visa_status", "Unknown")
    candidate.current_pipeline_stage = "Set 2 Passed"

    passed_s3, rej_stage_s3, rej_reason_s3, v_email, f_phone, website, portfolio = set3_contact.qualify_candidate_set3(candidate, profile)
    print(f"Test A (General Query Set 3): Passed={passed_s3}, Email='{v_email}', Phone='{f_phone}'")
    assert passed_s3 is True, f"Expected True for general query Set 3, got {passed_s3} ({rej_reason_s3})"

    candidate.email = v_email
    candidate.phone = f_phone
    candidate.website = website
    candidate.portfolio = portfolio
    candidate.current_pipeline_stage = "Qualified"

    # Test B: Clearance-Required Search ("Network Engineer TS/SCI")
    query_clearance = "Network Engineer TS/SCI"
    passed_s2_cleared, rej_stage_c, rej_reason_c, _, _ = set2_clearance.qualify_candidate_set2(candidate, profile, query_clearance)
    print(f"Test B (Clearance Required on Uncleared Candidate): Passed={passed_s2_cleared}, Stage='{rej_stage_c}', Reason='{rej_reason_c}'")
    assert passed_s2_cleared is False, f"Expected False for uncleared candidate on clearance query"

    # Test C: Cleared candidate on Clearance-Required Search
    profile_cleared = dict(profile)
    profile_cleared["bio"] = "Senior Network Engineer with active TS/SCI clearance and US Citizen."
    passed_s2_c2, _, _, _, clearance_c2 = set2_clearance.qualify_candidate_set2(candidate, profile_cleared, query_clearance)
    print(f"Test C (Clearance Required on Cleared Candidate): Passed={passed_s2_c2}, Clearance='{clearance_c2}', Visa='{profile_cleared.get('visa_status')}'")
    assert passed_s2_c2 is True, f"Expected True for cleared candidate on clearance query"
    assert clearance_c2 == "TS/SCI", f"Expected TS/SCI clearance, got {clearance_c2}"

    # Test D: Excel Export Verification
    test_excel_path = "scratch/test_profiles.xlsx"
    if os.path.exists(test_excel_path):
        os.remove(test_excel_path)
        
    set1_passed = [candidate]
    set2_passed = [candidate]
    set3_passed = [candidate]
    rejected_candidates = {"Set1": [], "Set2": [], "Set3": []}

    excel_manager.batch_save_profiles(set1_passed, set2_passed, set3_passed, rejected_candidates, test_excel_path)
    
    wb = openpyxl.load_workbook(test_excel_path)
    print("\nExcel Sheets Created:", wb.sheetnames)
    for sheet in ["SET_1", "SET_2", "SET_3"]:
        ws = wb[sheet]
        rows = list(ws.iter_rows(values_only=True))
        print(f"  {sheet}: Header={rows[0]}")
        print(f"  {sheet}: Row 1={rows[1]}")
        assert len(rows) == 2, f"Expected 2 rows in {sheet}, got {len(rows)}"

    print("\nALL QUALIFICATION & EXCEL EXPORT TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_tests()
