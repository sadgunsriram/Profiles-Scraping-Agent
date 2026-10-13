"""
sync_excel_to_db.py — Backfill existing ELIGIBLE candidates from Excel into PostgreSQL.
"""

import os
import sys
from datetime import datetime
from dotenv import load_dotenv
import openpyxl
from sqlalchemy.dialects.postgresql import insert as pg_insert

sys.path.insert(0, ".")
load_dotenv(".env")

from skills.profile_scraper.database_repository import (
    DatabaseCandidateRepository,
    EligibleCandidateModel
)

def sync_eligible_candidates():
    excel_path = os.getenv("EXCEL_OUTPUT_PATH", "skills/output/profiles.xlsx")
    if not os.path.exists(excel_path):
        print(f"Excel file not found at: {excel_path}")
        return

    repo = DatabaseCandidateRepository()
    if not repo.enabled or not repo.SessionLocal:
        print("[!] Database is not enabled or connection failed. Check .env settings.")
        return

    wb = openpyxl.load_workbook(excel_path)
    if "ELIGIBLE_CANDIDATES" not in wb.sheetnames:
        print("Sheet 'ELIGIBLE_CANDIDATES' not found in Excel workbook.")
        return

    ws = wb["ELIGIBLE_CANDIDATES"]
    headers = [str(cell).strip() if cell else "" for cell in next(ws.iter_rows(max_row=1, values_only=True))]
    print(f"Found sheet 'ELIGIBLE_CANDIDATES' with {ws.max_row - 1} data rows.")

    header_map = {name: idx for idx, name in enumerate(headers)}
    
    synced_count = 0
    with repo.SessionLocal() as session:
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or not any(row):
                continue
            
            def get_val(col_name, default="NOT_AVAILABLE"):
                idx = header_map.get(col_name)
                if idx is not None and idx < len(row):
                    val = row[idx]
                    if val is not None and str(val).strip() != "":
                        return val
                return default

            linkedin_url = str(get_val("LinkedIn URL", "")).strip()
            if not linkedin_url or linkedin_url == "NOT_AVAILABLE":
                continue

            full_name = str(get_val("Full Name", "Unknown Candidate")).strip()
            current_title = str(get_val("Current Role", "NOT_AVAILABLE")).strip()
            company = str(get_val("Company", "NOT_AVAILABLE")).strip()
            location = str(get_val("Location", "NOT_AVAILABLE")).strip()
            
            try:
                total_exp = float(get_val("Total Exp (Yrs)", 0.0))
            except (ValueError, TypeError):
                total_exp = 0.0

            try:
                relevant_tech_years = float(get_val("Relevant Tech Exp (Yrs)", 0.0))
            except (ValueError, TypeError):
                relevant_tech_years = 0.0

            try:
                skill_score = int(float(get_val("Skill Relevance Score", 0)))
            except (ValueError, TypeError):
                skill_score = 0

            exp_bracket = str(get_val("Experience Bracket", "NOT_AVAILABLE")).strip()
            primary_domain = str(get_val("Primary Domain", "NOT_AVAILABLE")).strip()
            tech_spec = str(get_val("Tech Specialization", "NOT_AVAILABLE")).strip()
            top_skills = str(get_val("Top Skills", "NOT_AVAILABLE")).strip()
            highest_edu = str(get_val("Highest Education", "NOT_AVAILABLE")).strip()
            work_auth = str(get_val("Work Authorization / Visa Status", "NOT_AVAILABLE")).strip()
            otw_raw = str(get_val("Open to Work", "YES")).strip().lower()
            is_otw = otw_raw.startswith("yes") or otw_raw == "true"

            email = str(get_val("Email", "")).strip()
            if email in ("NOT_AVAILABLE", "Not Available", ""):
                email = None

            phone = str(get_val("Phone", "")).strip()
            if phone in ("NOT_AVAILABLE", "Not Available", ""):
                phone = None

            sh_phone = str(get_val("SignalHire Verified Phone", "")).strip()
            if sh_phone in ("NOT_AVAILABLE", "Not Available", ""):
                sh_phone = None

            sh_email = str(get_val("SignalHire Verified Email", "")).strip()
            if sh_email in ("NOT_AVAILABLE", "Not Available", ""):
                sh_email = None

            sh_status = str(get_val("SignalHire Status", "NOT_AVAILABLE")).strip()
            github_url = str(get_val("GitHub URL", "")).strip()
            if github_url in ("NOT_AVAILABLE", "Not Available", ""):
                github_url = None

            raw_date = get_val("Date Qualified", None)
            date_qualified = datetime.utcnow()
            if isinstance(raw_date, datetime):
                date_qualified = raw_date
            elif isinstance(raw_date, str) and raw_date:
                try:
                    date_qualified = datetime.fromisoformat(raw_date)
                except Exception:
                    pass

            record_values = {
                "linkedin_url": linkedin_url,
                "full_name": full_name,
                "current_title": current_title,
                "company": company,
                "location": location,
                "total_experience_years": total_exp,
                "experience_bracket": exp_bracket,
                "relevant_tech_years": relevant_tech_years,
                "primary_domain": primary_domain,
                "tech_specialization": tech_spec,
                "skill_relevance_score": skill_score,
                "top_skills": top_skills,
                "highest_education": highest_edu,
                "work_authorization": work_auth,
                "is_open_to_work": is_otw,
                "email": email,
                "phone": phone,
                "signalhire_phone": sh_phone,
                "signalhire_email": sh_email,
                "signalhire_status": sh_status,
                "github_url": github_url,
                "date_qualified": date_qualified,
                "updated_at": datetime.utcnow(),
                "full_profile_json": "{}"
            }

            stmt = pg_insert(EligibleCandidateModel).values(**record_values)
            update_dict = {
                k: v for k, v in record_values.items()
                if k not in ("id", "linkedin_url", "date_qualified")
            }
            stmt = stmt.on_conflict_do_update(
                index_elements=["linkedin_url"],
                set_=update_dict
            )
            session.execute(stmt)
            synced_count += 1

        session.commit()

    print(f"Successfully synced {synced_count} eligible candidates from Excel to table 'eligible_candidates'!")
    total_db_count = repo.get_eligible_count()
    print(f"Total eligible candidates now in PostgreSQL: {total_db_count}")

if __name__ == "__main__":
    sync_eligible_candidates()
