"""
verify_complete_pipeline.py — End-to-end verification script for the decoupled 5-layer pipeline.
Verifies:
1. Complete progressive candidate profile normalization.
2. Candidate career analysis engine (timeline union, tech filtering, domain, explicit clearance/visa).
3. Eligibility evaluation into ELIGIBLE, NEEDS_REVIEW, and REJECTED.
4. Immediate 4-sheet Excel database persistence (ALL_RAW_PROFILES, ELIGIBLE_CANDIDATES, NEEDS_REVIEW, REJECTED_CANDIDATES).
"""

import os
import sys

# Ensure workspace root is in sys.path
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

import openpyxl
from skills.profile_scraper.candidate_models import (
    NormalizedCandidateProfile,
    ExperienceEntry,
    EducationEntry,
    CertificationEntry,
    ProjectEntry
)
from skills.profile_scraper.candidate_analyzer import CandidateAnalysisEngine
from skills.profile_scraper.eligibility_engine import EligibilityRulesEngine, JobRequirements
from skills.profile_scraper.excel_repository import ExcelCandidateRepository


def run_pipeline_verification():
    os.environ["LLM_ANALYZER_ENABLED"] = "false"
    excel_path = "skills/output/test_pipeline_profiles.xlsx"
    if os.path.exists(excel_path):
        try:
            os.remove(excel_path)
        except Exception:
            pass

    analyzer = CandidateAnalysisEngine()
    engine = EligibilityRulesEngine()
    repo = ExcelCandidateRepository(excel_path=excel_path)

    requirements = JobRequirements(
        target_role="Cloud Infrastructure Architect",
        min_relevant_experience_years=3.0,
        min_total_experience_years=4.0,
        require_us_location=True,
        required_security_clearance="Secret",
        reject_c2c_recruiters=True
    )

    print("\n" + "=" * 80)
    print("        DECOUPLED 5-LAYER PIPELINE END-TO-END VERIFICATION")
    print("=" * 80)

    # -------------------------------------------------------------
    # CANDIDATE 1: Strong Practitioner -> ELIGIBLE
    # -------------------------------------------------------------
    cand1 = NormalizedCandidateProfile(
        full_name="Marcus Vance",
        linkedin_url="https://www.linkedin.com/in/marcus-vance-cloud",
        headline="Senior Cloud Infrastructure Architect | AWS & Kubernetes",
        location="Austin, Texas, United States",
        about="AWS Architect with active Secret clearance. US Citizen specializing in Terraform and high-scale Kubernetes.",
        date_of_birth="NOT_AVAILABLE",
        gender="NOT_SPECIFIED",
        pronouns="He/Him",
        experience_history=[
            ExperienceEntry(
                title="Senior Cloud Architect",
                company="Nexus Cloud Solutions",
                start_date="Jan 2021",
                end_date="Present",
                description="Designed multi-region AWS infrastructure with Terraform and EKS."
            ),
            ExperienceEntry(
                title="DevOps Engineer",
                company="Austin Techworks",
                start_date="Jan 2018",
                end_date="Dec 2020",
                description="Automated CI/CD pipelines with Docker, Kubernetes, and Python."
            )
        ],
        education=[
            EducationEntry(institution="University of Texas at Austin", degree="Bachelor of Science", field_of_study="Computer Science", start_year=2014, end_year=2018)
        ],
        certifications=[
            CertificationEntry(name="AWS Certified Solutions Architect - Professional", issuing_organization="Amazon Web Services")
        ],
        projects=[
            ProjectEntry(title="Enterprise Multi-Cloud Mesh", description="Global Istio service mesh on EKS and GKE")
        ],
        skills=["AWS", "Terraform", "Kubernetes", "Docker", "Python", "Linux"],
        languages=["English - Native"],
        volunteer=["STEM Tutor | Austin Youth Coding"],
        email="marcus.vance@example.org",
        phone="+1 512-555-0199",
        websites=["https://marcusvance.dev"],
        github_url="https://github.com/marcusvance"
    )

    analysis1 = analyzer.analyze_candidate(cand1, target_role=requirements.target_role)
    decision1 = engine.evaluate(cand1, analysis1, requirements)
    repo.save_candidate_decision(decision1)
    print(f"[Profile 1] {cand1.full_name} -> DECISION: {decision1.decision.value}")
    print(f"            Total Exp: {analysis1.total_experience_years} yrs, Relevant Tech: {analysis1.relevant_technology_years} yrs, Domain: {analysis1.primary_professional_domain}")
    print(f"            Clearance: {analysis1.security_clearance}, Visa: {analysis1.work_authorization}, Score: {analysis1.skill_relevance_score}/100")

    # -------------------------------------------------------------
    # CANDIDATE 2: Strong Tech, but Unstated Clearance & No Direct Email -> NEEDS_REVIEW
    # -------------------------------------------------------------
    cand2 = NormalizedCandidateProfile(
        full_name="Samantha Wei",
        linkedin_url="https://www.linkedin.com/in/samantha-wei-devops",
        headline="Lead DevOps Engineer | Platform & Cloud Automation",
        location="Seattle, Washington, United States",
        about="Platform engineer building resilient cloud architectures. Open to exciting senior opportunities.",
        date_of_birth="NOT_AVAILABLE",
        gender="NOT_SPECIFIED",
        pronouns="She/Her",
        experience_history=[
            ExperienceEntry(
                title="Lead DevOps Engineer",
                company="PacWest Platforms",
                start_date="Mar 2020",
                end_date="Present",
                description="Built automated cloud infrastructure with AWS and Terraform."
            ),
            ExperienceEntry(
                title="Systems Engineer",
                company="Seattle Digital",
                start_date="Jun 2017",
                end_date="Feb 2020",
                description="Linux system administration and Docker containerization."
            )
        ],
        education=[
            EducationEntry(institution="University of Washington", degree="Bachelor of Science", field_of_study="Informatics", start_year=2013, end_year=2017)
        ],
        certifications=[
            CertificationEntry(name="Certified Kubernetes Administrator (CKA)", issuing_organization="Linux Foundation")
        ],
        skills=["AWS", "Kubernetes", "Terraform", "Go", "Docker"],
        email="NOT_AVAILABLE",
        phone="NOT_AVAILABLE",
        github_url="https://github.com/samanthawei"
    )

    analysis2 = analyzer.analyze_candidate(cand2, target_role=requirements.target_role)
    decision2 = engine.evaluate(cand2, analysis2, requirements)
    repo.save_candidate_decision(decision2)
    print(f"[Profile 2] {cand2.full_name} -> DECISION: {decision2.decision.value}")
    print(f"            Review Reasons: {decision2.review_reasons}")

    # -------------------------------------------------------------
    # CANDIDATE 3: Technical Recruiter -> REJECTED
    # -------------------------------------------------------------
    cand3 = NormalizedCandidateProfile(
        full_name="Gregory Hayes",
        linkedin_url="https://www.linkedin.com/in/gregory-hayes-recruiting",
        headline="Principal Talent Acquisition Partner | Technical Recruiting Specialist",
        location="Dallas, Texas, United States",
        about="Helping high-growth companies scale engineering organizations across Cloud, DevOps, and Full-Stack.",
        experience_history=[
            ExperienceEntry(
                title="Principal Technical Recruiter",
                company="Apex Talent Partners",
                start_date="Jan 2018",
                end_date="Present",
                description="Full-cycle recruiting for AWS, Kubernetes, and Python software engineers."
            )
        ],
        skills=["Technical Recruiting", "Talent Acquisition", "Sourcing", "Executive Search", "HR"],
        email="greg.hayes@example.org"
    )

    analysis3 = analyzer.analyze_candidate(cand3, target_role=requirements.target_role)
    decision3 = engine.evaluate(cand3, analysis3, requirements)
    repo.save_candidate_decision(decision3)
    print(f"[Profile 3] {cand3.full_name} -> DECISION: {decision3.decision.value}")
    print(f"            Failed Rules: {decision3.failed_mandatory_rules}")

    # -------------------------------------------------------------
    # VERIFY EXCEL WORKBOOK INTEGRITY
    # -------------------------------------------------------------
    wb = openpyxl.load_workbook(excel_path)
    print("\n" + "-" * 80)
    print(f"EXCEL DATABASE VERIFICATION: {excel_path}")
    print(f"Sheets in Workbook: {wb.sheetnames}")
    for name in ["ALL_RAW_PROFILES", "ELIGIBLE_CANDIDATES", "NEEDS_REVIEW", "REJECTED_CANDIDATES"]:
        ws = wb[name]
        print(f"  • Sheet '{name:<20}': {ws.max_row} rows (1 header + {ws.max_row - 1} data records)")

    assert wb["ALL_RAW_PROFILES"].max_row == 4  # 3 candidates + header
    assert wb["ELIGIBLE_CANDIDATES"].max_row == 2  # Marcus Vance
    assert wb["NEEDS_REVIEW"].max_row == 2  # Samantha Wei
    assert wb["REJECTED_CANDIDATES"].max_row == 2  # Gregory Hayes

    print("=" * 80)
    print("SUCCESS: Decoupled 5-layer pipeline and 4-sheet persistence verified perfectly!\n")


if __name__ == "__main__":
    run_pipeline_verification()
